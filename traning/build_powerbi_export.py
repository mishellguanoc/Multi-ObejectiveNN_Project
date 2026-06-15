# -*- coding: utf-8 -*-
"""
build_powerbi_export.py
================================================================
Lee el archivo TD_gloch_ec_v1.xlsx y genera un paquete completo
de migración a Power BI:

    powerbi_export/
        data/        -> CSVs (tablas de hechos + dimensiones + calendario)
        dax/         -> measures.dax (medidas DAX)
        tmdl/        -> modelo semántico TMDL (model.tmdl + tables/ + relaciones)
        dashboard/   -> dashboard_metadata.json (páginas y visuales)
        README.md    -> instrucciones de uso

Flujo:  Python -> Genera datos -> Genera DAX -> Genera TMDL -> Genera metadatos
================================================================
"""

import os
import csv
import json
import uuid
import datetime as dt
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

# ─────────────────────────────────────────────────────────────
# RUTAS
# ─────────────────────────────────────────────────────────────
EXCEL_NAME = "TD_gloch_ec_v1.xlsx"


def _locate_excel():
    """Busca el Excel en ubicaciones candidatas (carpeta del script,
    carpeta padre, directorio actual)."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        os.path.join(script_dir, EXCEL_NAME),
        os.path.join(os.path.dirname(script_dir), EXCEL_NAME),
        os.path.join(os.getcwd(), EXCEL_NAME),
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    # Búsqueda recursiva como último recurso
    for base in (script_dir, os.path.dirname(script_dir)):
        for root, _dirs, files in os.walk(base):
            if EXCEL_NAME in files:
                return os.path.join(root, EXCEL_NAME)
    return candidates[0]  # devuelve la ruta por defecto (fallará con mensaje claro)


EXCEL_PATH = _locate_excel()
BASE_DIR   = os.path.dirname(EXCEL_PATH)          # salida junto al Excel
OUT_DIR    = os.path.join(BASE_DIR, "powerbi_export")
DATA_DIR   = os.path.join(OUT_DIR, "data")
DAX_DIR    = os.path.join(OUT_DIR, "dax")
TMDL_DIR   = os.path.join(OUT_DIR, "tmdl")
TABLES_DIR = os.path.join(TMDL_DIR, "tables")
DASH_DIR   = os.path.join(OUT_DIR, "dashboard")

TODAY = dt.date.today()

MESES_ES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

# Mapa hoja -> nombre de tabla estructurada
TABLE_SHEETS = {
    "CONTROL":       "tbl_CONTROL",
    "MADUREZ":       "tbl_MADUREZ",
    "ANTES_DESPUES": "tbl_BRECHAS",
    "RIESGOS":       "tbl_RIESGOS",
    "BACKLOG":       "tbl_BACKLOG",
    "KPIs":          "tbl_KPIs",
}

# Columnas que en Excel son fórmulas: se recalculan en Python
DERIVED_COLUMNS = {
    "tbl_CONTROL": ["Duracion_Dias", "Semanas_Rest", "Alerta_Vencim"],
    "tbl_MADUREZ": ["Brecha", "Punt_Ponderada"],
    "tbl_BRECHAS": ["Semaforo"],
    "tbl_RIESGOS": ["Criticidad", "Nivel_Criticidad"],
    "tbl_BACKLOG": ["Score_Prioridad", "Clasificacion_MoSCoW"],
    "tbl_KPIs":    ["Valor_Actual", "Semaforo"],
}


# ═════════════════════════════════════════════════════════════
# 1) LECTURA DEL EXCEL
# ═════════════════════════════════════════════════════════════
def read_tables():
    print("  Leyendo tablas del Excel...")
    wb = load_workbook(EXCEL_PATH, data_only=False)
    tables = {}

    for sheet_name, table_name in TABLE_SHEETS.items():
        ws = wb[sheet_name]
        if table_name not in ws.tables:
            raise RuntimeError(f"No se encontró la tabla {table_name} en {sheet_name}")
        ref = ws.tables[table_name].ref
        min_col, min_row, max_col, max_row = range_boundaries(ref)

        # Encabezados
        headers = []
        for c in range(min_col, max_col + 1):
            headers.append(str(ws.cell(row=min_row, column=c).value).strip())

        # Filas de datos
        rows = []
        for r in range(min_row + 1, max_row + 1):
            row = {}
            for idx, c in enumerate(range(min_col, max_col + 1)):
                val = ws.cell(row=r, column=c).value
                # Las fórmulas se descartan (se recalculan después)
                if isinstance(val, str) and val.startswith("="):
                    val = None
                row[headers[idx]] = val
            rows.append(row)

        tables[table_name] = {"headers": headers, "rows": rows}
        print(f"    - {table_name}: {len(rows)} filas, {len(headers)} columnas")

    return tables


# ═════════════════════════════════════════════════════════════
# 2) RECÁLCULO DE COLUMNAS DERIVADAS
# ═════════════════════════════════════════════════════════════
def _to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def recompute_derived(tables):
    print("  Recalculando columnas derivadas...")

    # ── CONTROL
    for row in tables["tbl_CONTROL"]["rows"]:
        fi  = _to_date(row.get("Fecha_Inicio"))
        ffp = _to_date(row.get("Fecha_Fin_Plan"))
        ffr = _to_date(row.get("Fecha_Fin_Real"))
        estado = row.get("Estado") or ""
        # Duración
        ref_fin = ffr if ffr else ffp
        row["Duracion_Dias"] = (ref_fin - fi).days if (fi and ref_fin) else None
        # Semanas restantes
        if ffp:
            diff = (ffp - TODAY).days
            row["Semanas_Rest"] = max(0, -(-diff // 7)) if diff > 0 else 0
        else:
            row["Semanas_Rest"] = None
        # Alerta de vencimiento
        if ffp:
            if ffp < TODAY and estado != "Completado":
                row["Alerta_Vencim"] = "VENCIDO"
            elif (ffp - TODAY).days <= 7:
                row["Alerta_Vencim"] = "PROXIMO"
            else:
                row["Alerta_Vencim"] = "OK"
        else:
            row["Alerta_Vencim"] = "OK"

    # ── MADUREZ
    for row in tables["tbl_MADUREZ"]["rows"]:
        na = row.get("Nivel_Actual") or 0
        no = row.get("Nivel_Objetivo") or 0
        peso = row.get("Peso_Pct") or 0
        row["Brecha"] = no - na
        row["Punt_Ponderada"] = round(na * peso, 4)

    # ── BRECHAS
    for row in tables["tbl_BRECHAS"]["rows"]:
        pct = row.get("Pct_Cumplimiento") or 0
        if pct >= 0.7:
            row["Semaforo"] = "VERDE"
        elif pct >= 0.3:
            row["Semaforo"] = "AMARILLO"
        else:
            row["Semaforo"] = "ROJO"

    # ── RIESGOS
    for row in tables["tbl_RIESGOS"]["rows"]:
        p = row.get("Probabilidad") or 0
        i = row.get("Impacto") or 0
        crit = p * i
        row["Criticidad"] = crit
        if crit >= 15:
            row["Nivel_Criticidad"] = "CRITICO"
        elif crit >= 9:
            row["Nivel_Criticidad"] = "ALTO"
        elif crit >= 4:
            row["Nivel_Criticidad"] = "MEDIO"
        else:
            row["Nivel_Criticidad"] = "BAJO"

    # ── BACKLOG
    for row in tables["tbl_BACKLOG"]["rows"]:
        v = row.get("Valor_Negocio") or 0
        u = row.get("Urgencia") or 0
        c = row.get("Complejidad") or 1
        score = round((v * 2 + u * 2) / c, 2) if c else 0
        row["Score_Prioridad"] = score
        if score >= 6:
            row["Clasificacion_MoSCoW"] = "Must Have"
        elif score >= 4:
            row["Clasificacion_MoSCoW"] = "Should Have"
        elif score >= 2:
            row["Clasificacion_MoSCoW"] = "Could Have"
        else:
            row["Clasificacion_MoSCoW"] = "Won't Have"

    # ── KPIs (Valor_Actual calculado desde las demás tablas)
    ctrl = tables["tbl_CONTROL"]["rows"]
    rsk  = tables["tbl_RIESGOS"]["rows"]
    mad  = tables["tbl_MADUREZ"]["rows"]

    def avg_avance(area=None):
        vals = [r["Pct_Avance"] for r in ctrl
                if r.get("Pct_Avance") is not None and (area is None or r.get("Área_Estratégica") == area)]
        return round(sum(vals) / len(vals), 4) if vals else 0

    completados = sum(1 for r in ctrl if r.get("Estado") == "Completado")
    bloqueadas  = sum(1 for r in ctrl if r.get("Estado") == "Bloqueado")
    riesgos_crit_act = sum(1 for r in rsk
                           if r.get("Nivel_Criticidad") == "CRITICO" and r.get("Estado") != "Cerrado")
    riesgos_tot = len(rsk)
    dur_comp = [r["Duracion_Dias"] for r in ctrl
                if r.get("Estado") == "Completado" and r.get("Duracion_Dias") is not None]
    tiempo_prom = round(sum(dur_comp) / len(dur_comp), 1) if dur_comp else 0
    madurez = round(sum((r.get("Nivel_Actual") or 0) * (r.get("Peso_Pct") or 0) for r in mad) /
                    (sum((r.get("Peso_Pct") or 0) for r in mad) or 1), 2)

    kpi_values = {
        "KPI-001": avg_avance(),
        "KPI-002": madurez,
        "KPI-003": completados,
        "KPI-004": riesgos_crit_act,
        "KPI-005": bloqueadas,
        "KPI-006": tiempo_prom,
        "KPI-007": avg_avance("Infraestructura TI"),
        "KPI-008": avg_avance("Seguridad Digital"),
        "KPI-009": avg_avance("Gestión IT"),
        "KPI-010": avg_avance("Presencia Digital"),
        "KPI-011": 0.05,   # manual
        "KPI-012": riesgos_tot,
    }

    for row in tables["tbl_KPIs"]["rows"]:
        kid = row.get("ID_KPI")
        if kid in kpi_values:
            row["Valor_Actual"] = kpi_values[kid]
        meta = row.get("Meta")
        va = row.get("Valor_Actual")
        tend = row.get("Tendencia") or "↑"
        # Semáforo
        try:
            if tend == "↑":
                row["Semaforo"] = "VERDE" if va >= meta else ("AMARILLO" if va >= meta * 0.7 else "ROJO")
            else:
                row["Semaforo"] = "VERDE" if va <= meta else ("AMARILLO" if va <= meta * 1.3 else "ROJO")
        except (TypeError, ZeroDivisionError):
            row["Semaforo"] = "AMARILLO"

    return tables


# ═════════════════════════════════════════════════════════════
# 3) DIMENSIONES + CALENDARIO
# ═════════════════════════════════════════════════════════════
def build_dimensions(tables):
    dims = {}
    dims["Dim_Area"] = {
        "headers": ["Area"],
        "rows": [{"Area": a} for a in [
            "Infraestructura TI", "Seguridad Digital", "Presencia Digital",
            "Gestión IT", "Gestión de Datos", "Conectividad", "Automatización"]]
    }
    dims["Dim_Estado"] = {
        "headers": ["Estado", "Orden"],
        "rows": [
            {"Estado": "No iniciado", "Orden": 1},
            {"Estado": "En progreso", "Orden": 2},
            {"Estado": "Bloqueado",   "Orden": 3},
            {"Estado": "Completado",  "Orden": 4},
        ]
    }
    dims["Dim_Prioridad"] = {
        "headers": ["Prioridad", "Orden"],
        "rows": [
            {"Prioridad": "Alta",  "Orden": 1},
            {"Prioridad": "Media", "Orden": 2},
            {"Prioridad": "Baja",  "Orden": 3},
        ]
    }
    dims["Dim_Categoria_Riesgo"] = {
        "headers": ["Categoria"],
        "rows": [{"Categoria": c} for c in
                 ["Tecnológico", "Operativo", "Seguridad", "Continuidad", "Proyecto"]]
    }
    dims["Dim_Fase"] = {
        "headers": ["Fase", "Orden"],
        "rows": [
            {"Fase": "Fase 1", "Orden": 1},
            {"Fase": "Fase 2", "Orden": 2},
            {"Fase": "Fase 3", "Orden": 3},
        ]
    }

    # ── Dim_Calendar (rango basado en fechas del CONTROL)
    all_dates = []
    for r in tables["tbl_CONTROL"]["rows"]:
        for col in ("Fecha_Inicio", "Fecha_Fin_Plan", "Fecha_Fin_Real"):
            d = _to_date(r.get(col))
            if d:
                all_dates.append(d)
    if all_dates:
        dmin, dmax = min(all_dates), max(all_dates)
    else:
        dmin, dmax = TODAY, TODAY
    # Expandir a inicio/fin de año
    start = dt.date(dmin.year, 1, 1)
    end   = dt.date(dmax.year, 12, 31)

    cal_rows = []
    cur = start
    while cur <= end:
        cal_rows.append({
            "Date": cur.isoformat(),
            "Anio": cur.year,
            "Mes": cur.month,
            "NombreMes": MESES_ES[cur.month],
            "Trimestre": "T{}".format((cur.month - 1) // 3 + 1),
            "AnioMes": "{:04d}-{:02d}".format(cur.year, cur.month),
            "Semana": cur.isocalendar()[1],
            "DiaSemana": cur.isoweekday(),
        })
        cur += dt.timedelta(days=1)

    dims["Dim_Calendar"] = {
        "headers": ["Date", "Anio", "Mes", "NombreMes", "Trimestre", "AnioMes", "Semana", "DiaSemana"],
        "rows": cal_rows
    }
    return dims


# ═════════════════════════════════════════════════════════════
# ESQUEMA DE TIPOS  (col -> (tmdl_type, m_type, formatString))
# ═════════════════════════════════════════════════════════════
SCHEMA = {
    # CONTROL
    "Pct_Avance": ("double", "type number", "0.0%"),
    "Fecha_Inicio": ("dateTime", "type date", "dd/mm/yyyy"),
    "Fecha_Fin_Plan": ("dateTime", "type date", "dd/mm/yyyy"),
    "Fecha_Fin_Real": ("dateTime", "type date", "dd/mm/yyyy"),
    "Duracion_Dias": ("int64", "Int64.Type", "0"),
    "Semanas_Rest": ("int64", "Int64.Type", "0"),
    # MADUREZ
    "Nivel_Actual": ("int64", "Int64.Type", "0"),
    "Nivel_Objetivo": ("int64", "Int64.Type", "0"),
    "Brecha": ("int64", "Int64.Type", "0"),
    "Peso_Pct": ("double", "type number", "0%"),
    "Punt_Ponderada": ("double", "type number", "0.00"),
    # BRECHAS
    "Pct_Cumplimiento": ("double", "type number", "0%"),
    # RIESGOS
    "Probabilidad": ("int64", "Int64.Type", "0"),
    "Impacto": ("int64", "Int64.Type", "0"),
    "Criticidad": ("int64", "Int64.Type", "0"),
    "Fecha_Identificacion": ("dateTime", "type date", "dd/mm/yyyy"),
    "Fecha_Revision": ("dateTime", "type date", "dd/mm/yyyy"),
    # BACKLOG
    "Valor_Negocio": ("int64", "Int64.Type", "0"),
    "Urgencia": ("int64", "Int64.Type", "0"),
    "Complejidad": ("int64", "Int64.Type", "0"),
    "Score_Prioridad": ("double", "type number", "0.0"),
    # KPIs
    "Valor_Actual": ("double", "type number", None),
    "Meta": ("double", "type number", None),
    # Dimensiones
    "Orden": ("int64", "Int64.Type", "0"),
    "Anio": ("int64", "Int64.Type", "0"),
    "Mes": ("int64", "Int64.Type", "0"),
    "Semana": ("int64", "Int64.Type", "0"),
    "DiaSemana": ("int64", "Int64.Type", "0"),
    "Date": ("dateTime", "type date", "dd/mm/yyyy"),
}


def col_type(col):
    return SCHEMA.get(col, ("string", "type text", None))


# ═════════════════════════════════════════════════════════════
# 4) ESCRITURA DE CSV
# ═════════════════════════════════════════════════════════════
def fmt_value(col, v):
    if v is None:
        return ""
    d = _to_date(v) if isinstance(v, (dt.date, dt.datetime)) else None
    if d:
        return d.isoformat()
    return v


def write_csvs(all_tables):
    print("  Generando CSVs en data/ ...")
    for name, tbl in all_tables.items():
        path = os.path.join(DATA_DIR, f"{name}.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(tbl["headers"])
            for row in tbl["rows"]:
                w.writerow([fmt_value(h, row.get(h)) for h in tbl["headers"]])
        print(f"    - {name}.csv ({len(tbl['rows'])} filas)")


# ═════════════════════════════════════════════════════════════
# 5) MEDIDAS DAX
# ═════════════════════════════════════════════════════════════
# (nombre, expresión, formatString)
MEASURES = [
    # Avance
    ("Total Iniciativas", "COUNTROWS ( tbl_CONTROL )", "0"),
    ("Iniciativas Completadas", 'CALCULATE ( COUNTROWS ( tbl_CONTROL ), tbl_CONTROL[Estado] = "Completado" )', "0"),
    ("Iniciativas En Progreso", 'CALCULATE ( COUNTROWS ( tbl_CONTROL ), tbl_CONTROL[Estado] = "En progreso" )', "0"),
    ("Iniciativas Bloqueadas", 'CALCULATE ( COUNTROWS ( tbl_CONTROL ), tbl_CONTROL[Estado] = "Bloqueado" )', "0"),
    ("Iniciativas No Iniciadas", 'CALCULATE ( COUNTROWS ( tbl_CONTROL ), tbl_CONTROL[Estado] = "No iniciado" )', "0"),
    ("Iniciativas Vencidas", 'CALCULATE ( COUNTROWS ( tbl_CONTROL ), tbl_CONTROL[Alerta_Vencim] = "VENCIDO" )', "0"),
    ("Avance Global", "AVERAGE ( tbl_CONTROL[Pct_Avance] )", "0.0%"),
    ("% Completado", "DIVIDE ( [Iniciativas Completadas], [Total Iniciativas] )", "0.0%"),
    # Riesgos
    ("Total Riesgos", "COUNTROWS ( tbl_RIESGOS )", "0"),
    ("Riesgos Criticos", 'CALCULATE ( COUNTROWS ( tbl_RIESGOS ), tbl_RIESGOS[Nivel_Criticidad] = "CRITICO" )', "0"),
    ("Riesgos Criticos Activos",
     'CALCULATE ( COUNTROWS ( tbl_RIESGOS ), tbl_RIESGOS[Nivel_Criticidad] = "CRITICO", tbl_RIESGOS[Estado] <> "Cerrado" )', "0"),
    ("Riesgos Abiertos", 'CALCULATE ( COUNTROWS ( tbl_RIESGOS ), tbl_RIESGOS[Estado] <> "Cerrado" )', "0"),
    ("Criticidad Promedio", "AVERAGE ( tbl_RIESGOS[Criticidad] )", "0.0"),
    ("Criticidad Maxima", "MAX ( tbl_RIESGOS[Criticidad] )", "0"),
    # Madurez
    ("Madurez Global",
     "DIVIDE ( SUMX ( tbl_MADUREZ, tbl_MADUREZ[Nivel_Actual] * tbl_MADUREZ[Peso_Pct] ), SUM ( tbl_MADUREZ[Peso_Pct] ) )", "0.00"),
    ("Madurez Objetivo",
     "DIVIDE ( SUMX ( tbl_MADUREZ, tbl_MADUREZ[Nivel_Objetivo] * tbl_MADUREZ[Peso_Pct] ), SUM ( tbl_MADUREZ[Peso_Pct] ) )", "0.00"),
    ("Brecha Madurez", "[Madurez Objetivo] - [Madurez Global]", "0.00"),
    ("Nivel Madurez Texto",
     'SWITCH ( TRUE (), [Madurez Global] >= 4, "Gestionado", [Madurez Global] >= 3, "Definido", [Madurez Global] >= 2, "Repetible", "Inicial" )', None),
    # Brechas
    ("Cumplimiento Promedio", "AVERAGE ( tbl_BRECHAS[Pct_Cumplimiento] )", "0.0%"),
    ("Brechas Criticas", 'CALCULATE ( COUNTROWS ( tbl_BRECHAS ), tbl_BRECHAS[Nivel_Brecha] = "Alta" )', "0"),
    # Backlog
    ("Total Backlog", "COUNTROWS ( tbl_BACKLOG )", "0"),
    ("Backlog Must Have", 'CALCULATE ( COUNTROWS ( tbl_BACKLOG ), tbl_BACKLOG[Clasificacion_MoSCoW] = "Must Have" )', "0"),
    ("Score Promedio Backlog", "AVERAGE ( tbl_BACKLOG[Score_Prioridad] )", "0.0"),
    ("Backlog Fase 1", 'CALCULATE ( COUNTROWS ( tbl_BACKLOG ), tbl_BACKLOG[Fase_Propuesta] = "Fase 1" )', "0"),
]

MEASURES_TABLE = "tbl_KPIs"  # tabla anfitriona de las medidas


def write_dax():
    print("  Generando dax/measures.dax ...")
    path = os.path.join(DAX_DIR, "measures.dax")
    with open(path, "w", encoding="utf-8") as f:
        f.write("// ================================================================\n")
        f.write("// MEDIDAS DAX — Sistema de Transformación Digital gloch.ec\n")
        f.write("// Tabla anfitriona sugerida: " + MEASURES_TABLE + "\n")
        f.write("// Copie cada medida en Power BI: Modelado > Nueva medida\n")
        f.write("// ================================================================\n\n")
        for name, expr, fmt in MEASURES:
            if fmt:
                f.write(f"// formato: {fmt}\n")
            f.write(f"{name} = {expr}\n\n")
    print(f"    - measures.dax ({len(MEASURES)} medidas)")


# ═════════════════════════════════════════════════════════════
# 6) MODELO TMDL
# ═════════════════════════════════════════════════════════════
RELATIONSHIPS = [
    ("rel_control_area",      "tbl_CONTROL", "Área_Estratégica", "Dim_Area", "Area"),
    ("rel_control_estado",    "tbl_CONTROL", "Estado",           "Dim_Estado", "Estado"),
    ("rel_control_prioridad", "tbl_CONTROL", "Prioridad",        "Dim_Prioridad", "Prioridad"),
    ("rel_control_riesgo",    "tbl_CONTROL", "Riesgo_Asociado",  "tbl_RIESGOS", "ID_Riesgo"),
    ("rel_riesgo_categoria",  "tbl_RIESGOS", "Categoria",        "Dim_Categoria_Riesgo", "Categoria"),
    ("rel_backlog_fase",      "tbl_BACKLOG", "Fase_Propuesta",   "Dim_Fase", "Fase"),
    ("rel_control_calendar",  "tbl_CONTROL", "Fecha_Inicio",     "Dim_Calendar", "Date"),
]


def needs_quote(name):
    return not name.replace("_", "").isalnum() or " " in name or not name.isascii()


def q(name):
    """Cita el nombre del objeto en TMDL si es necesario."""
    return f"'{name}'" if needs_quote(name) else name


def m_path(name):
    return os.path.join(DATA_DIR, f"{name}.csv").replace("\\", "\\")


def tmdl_partition(name, headers):
    """Genera la partición Power Query (M) que lee el CSV."""
    csv_path = os.path.join(DATA_DIR, f"{name}.csv")
    transforms = []
    for h in headers:
        _, mt, _ = col_type(h)
        transforms.append(f'{{"{h}", {mt}}}')
    transforms_str = ", ".join(transforms)
    lines = []
    lines.append(f"\tpartition {q(name)} = m")
    lines.append("\t\tmode: import")
    lines.append("\t\tsource =")
    lines.append("\t\t\t\tlet")
    lines.append(f'\t\t\t\t\tOrigen = Csv.Document(File.Contents("{csv_path}"), [Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),')
    lines.append("\t\t\t\t\tEncabezados = Table.PromoteHeaders(Origen, [PromoteAllScalars=true]),")
    lines.append(f"\t\t\t\t\tTipos = Table.TransformColumnTypes(Encabezados, {{{transforms_str}}})")
    lines.append("\t\t\t\tin")
    lines.append("\t\t\t\t\tTipos")
    return "\n".join(lines)


def write_table_tmdl(name, headers, measures=None):
    lines = [f"table {q(name)}", ""]
    # Columnas
    for h in headers:
        tt, _, fmt = col_type(h)
        lines.append(f"\tcolumn {q(h)}")
        lines.append(f"\t\tdataType: {tt}")
        if tt == "dateTime":
            lines.append("\t\tformatString: General Date")
        elif fmt:
            lines.append(f"\t\tformatString: {fmt}")
        lines.append("\t\tsummarizeBy: none")
        lines.append(f"\t\tsourceColumn: {h}")
        lines.append("")
    # Medidas (solo en la tabla anfitriona)
    if measures:
        for mname, expr, fmt in measures:
            lines.append(f"\tmeasure {q(mname)} = {expr}")
            if fmt:
                lines.append(f"\t\tformatString: {fmt}")
            lines.append("")
    # Partición
    lines.append(tmdl_partition(name, headers))
    lines.append("")
    path = os.path.join(TABLES_DIR, f"{name}.tmdl")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def write_model_tmdl(all_tables):
    lines = ["model Model",
             "\tculture: es-EC",
             "\tdefaultPowerBIDataSourceVersion: powerBI_V3",
             "\tsourceQueryCulture: es-EC",
             ""]
    for name in all_tables.keys():
        lines.append(f"\tref table {q(name)}")
    lines.append("")
    # Relaciones
    for rid, ft, fc, tt, tc in RELATIONSHIPS:
        lines.append(f"relationship {rid}")
        lines.append(f"\tfromColumn: {ft}.{q(fc)}")
        lines.append(f"\ttoColumn: {tt}.{q(tc)}")
        lines.append("\tcrossFilteringBehavior: oneDirection")
        lines.append("")
    with open(os.path.join(TMDL_DIR, "model.tmdl"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def write_tmdl(all_tables):
    print("  Generando modelo TMDL en tmdl/ ...")
    for name, tbl in all_tables.items():
        measures = MEASURES if name == MEASURES_TABLE else None
        write_table_tmdl(name, tbl["headers"], measures)
    write_model_tmdl(all_tables)
    print(f"    - model.tmdl + {len(all_tables)} tablas + {len(RELATIONSHIPS)} relaciones")


# ═════════════════════════════════════════════════════════════
# 7) METADATOS DEL DASHBOARD
# ═════════════════════════════════════════════════════════════
def write_dashboard_metadata():
    print("  Generando dashboard/dashboard_metadata.json ...")
    meta = {
        "report_name": "Transformación Digital gloch.ec",
        "version": "1.0",
        "locale": "es-EC",
        "generated_at": TODAY.isoformat(),
        "theme": {
            "primary": "#1F3864",
            "secondary": "#2E75B6",
            "accent": "#FFC000",
            "positive": "#1E7E34",
            "warning": "#FF8C00",
            "negative": "#C0392B",
            "background": "#FFFFFF",
            "font_family": "Segoe UI"
        },
        "model": {
            "fact_tables": ["tbl_CONTROL", "tbl_RIESGOS", "tbl_KPIs"],
            "dimension_tables": ["Dim_Area", "Dim_Estado", "Dim_Prioridad",
                                  "Dim_Categoria_Riesgo", "Dim_Fase", "Dim_Calendar",
                                  "tbl_MADUREZ", "tbl_BRECHAS", "tbl_BACKLOG"],
            "measures_table": MEASURES_TABLE
        },
        "pages": [
            {
                "name": "Resumen Ejecutivo",
                "order": 1,
                "layout": "grid-12",
                "visuals": [
                    {"id": "kpi_avance", "type": "card", "title": "% Avance Global",
                     "field": "[Avance Global]", "position": {"x": 0, "y": 0, "w": 2, "h": 2}},
                    {"id": "kpi_completados", "type": "card", "title": "Iniciativas Completadas",
                     "field": "[Iniciativas Completadas]", "position": {"x": 2, "y": 0, "w": 2, "h": 2}},
                    {"id": "kpi_riesgos", "type": "card", "title": "Riesgos Críticos Activos",
                     "field": "[Riesgos Criticos Activos]", "position": {"x": 4, "y": 0, "w": 2, "h": 2}},
                    {"id": "kpi_bloqueadas", "type": "card", "title": "Tareas Bloqueadas",
                     "field": "[Iniciativas Bloqueadas]", "position": {"x": 6, "y": 0, "w": 2, "h": 2}},
                    {"id": "gauge_madurez", "type": "gauge", "title": "Madurez Digital",
                     "value": "[Madurez Global]", "target": "[Madurez Objetivo]", "min": 0, "max": 5,
                     "position": {"x": 8, "y": 0, "w": 4, "h": 2}},
                    {"id": "bar_avance_area", "type": "clusteredBarChart", "title": "Avance por Área",
                     "axis": "Dim_Area[Area]", "values": ["[Avance Global]"],
                     "position": {"x": 0, "y": 2, "w": 6, "h": 4}},
                    {"id": "donut_estados", "type": "donutChart", "title": "Distribución de Estados",
                     "legend": "Dim_Estado[Estado]", "values": ["[Total Iniciativas]"],
                     "position": {"x": 6, "y": 2, "w": 6, "h": 4}},
                    {"id": "slicer_area", "type": "slicer", "title": "Área",
                     "field": "Dim_Area[Area]", "position": {"x": 0, "y": 6, "w": 3, "h": 2}},
                    {"id": "slicer_fecha", "type": "slicer", "title": "Periodo",
                     "field": "Dim_Calendar[AnioMes]", "position": {"x": 3, "y": 6, "w": 3, "h": 2}},
                ]
            },
            {
                "name": "Madurez Digital",
                "order": 2,
                "visuals": [
                    {"id": "radar_madurez", "type": "lineChart", "title": "Madurez Actual vs Objetivo",
                     "axis": "tbl_MADUREZ[Dimensión]",
                     "values": ["tbl_MADUREZ[Nivel_Actual]", "tbl_MADUREZ[Nivel_Objetivo]"],
                     "note": "Usar visual de radar (marketplace) o gráfico de líneas",
                     "position": {"x": 0, "y": 0, "w": 6, "h": 5}},
                    {"id": "table_madurez", "type": "table", "title": "Detalle de Madurez",
                     "columns": ["tbl_MADUREZ[Dimensión]", "tbl_MADUREZ[Nivel_Actual]",
                                 "tbl_MADUREZ[Nivel_Objetivo]", "tbl_MADUREZ[Brecha]", "tbl_MADUREZ[Peso_Pct]"],
                     "position": {"x": 6, "y": 0, "w": 6, "h": 5}},
                    {"id": "card_madurez_global", "type": "card", "title": "Madurez Global",
                     "field": "[Madurez Global]", "position": {"x": 0, "y": 5, "w": 3, "h": 2}},
                    {"id": "card_brecha", "type": "card", "title": "Brecha de Madurez",
                     "field": "[Brecha Madurez]", "position": {"x": 3, "y": 5, "w": 3, "h": 2}},
                ]
            },
            {
                "name": "Riesgos",
                "order": 3,
                "visuals": [
                    {"id": "matrix_riesgos", "type": "matrix", "title": "Matriz de Criticidad",
                     "rows": "tbl_RIESGOS[Probabilidad]", "columns": "tbl_RIESGOS[Impacto]",
                     "values": ["[Total Riesgos]"], "position": {"x": 0, "y": 0, "w": 5, "h": 5}},
                    {"id": "bar_riesgo_cat", "type": "clusteredColumnChart", "title": "Riesgos por Categoría",
                     "axis": "Dim_Categoria_Riesgo[Categoria]", "values": ["[Total Riesgos]"],
                     "position": {"x": 5, "y": 0, "w": 7, "h": 5}},
                    {"id": "table_riesgos", "type": "table", "title": "Detalle de Riesgos",
                     "columns": ["tbl_RIESGOS[ID_Riesgo]", "tbl_RIESGOS[Descripcion_Riesgo]",
                                 "tbl_RIESGOS[Nivel_Criticidad]", "tbl_RIESGOS[Estado]",
                                 "tbl_RIESGOS[Plan_Mitigacion]"],
                     "conditional_formatting": {"field": "tbl_RIESGOS[Nivel_Criticidad]",
                                                  "rules": {"CRITICO": "#C0392B", "ALTO": "#FF8C00",
                                                            "MEDIO": "#FFC000", "BAJO": "#1E7E34"}},
                     "position": {"x": 0, "y": 5, "w": 12, "h": 4}},
                ]
            },
            {
                "name": "Backlog y Hoja de Ruta",
                "order": 4,
                "visuals": [
                    {"id": "bar_moscow", "type": "clusteredBarChart", "title": "Backlog por Clasificación MoSCoW",
                     "axis": "tbl_BACKLOG[Clasificacion_MoSCoW]", "values": ["[Total Backlog]"],
                     "position": {"x": 0, "y": 0, "w": 6, "h": 4}},
                    {"id": "bar_fase", "type": "stackedColumnChart", "title": "Iniciativas por Fase",
                     "axis": "Dim_Fase[Fase]", "legend": "tbl_BACKLOG[Tema_Principal]",
                     "values": ["[Total Backlog]"], "position": {"x": 6, "y": 0, "w": 6, "h": 4}},
                    {"id": "table_backlog", "type": "table", "title": "Backlog Priorizado",
                     "columns": ["tbl_BACKLOG[Iniciativa_Descripcion]", "tbl_BACKLOG[Score_Prioridad]",
                                 "tbl_BACKLOG[Clasificacion_MoSCoW]", "tbl_BACKLOG[Fase_Propuesta]",
                                 "tbl_BACKLOG[Estado]"],
                     "sort": {"by": "tbl_BACKLOG[Score_Prioridad]", "direction": "desc"},
                     "position": {"x": 0, "y": 4, "w": 12, "h": 5}},
                ]
            }
        ]
    }
    with open(os.path.join(DASH_DIR, "dashboard_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    print(f"    - dashboard_metadata.json ({len(meta['pages'])} páginas)")


# ═════════════════════════════════════════════════════════════
# 8) README
# ═════════════════════════════════════════════════════════════
def write_readme(all_tables):
    print("  Generando README.md ...")
    content = f"""# Paquete de migración a Power BI — gloch.ec

Generado automáticamente desde `TD_gloch_ec_v1.xlsx` el {TODAY.isoformat()}.

## Estructura

```
powerbi_export/
├── data/         CSVs listos para importar (UTF-8 con BOM)
├── dax/          measures.dax (medidas DAX)
├── tmdl/         Modelo semántico TMDL (model.tmdl + tables/)
├── dashboard/    dashboard_metadata.json (diseño de páginas y visuales)
└── README.md
```

## Tablas generadas

| Tabla | Tipo | Filas |
|-------|------|-------|
""" + "\n".join(
        f"| {n} | {'Hecho' if n in ('tbl_CONTROL','tbl_RIESGOS','tbl_KPIs') else 'Dimensión/Soporte'} | {len(t['rows'])} |"
        for n, t in all_tables.items()
    ) + f"""

## Opción A — Importación rápida (recomendada para empezar)

1. Abre **Power BI Desktop**.
2. `Obtener datos` > `Texto/CSV` e importa cada archivo de `data/`.
   - Verifica tipos: fechas como Fecha, `Pct_*` y `Peso_Pct` como Número decimal, niveles como Número entero.
3. En la vista **Modelo**, crea las relaciones:
"""
    for rid, ft, fc, tt, tc in RELATIONSHIPS:
        content += f"   - `{ft}[{fc}]` (*) -> `{tt}[{tc}]` (1)\n"
    content += """4. Marca `Dim_Calendar` como **tabla de fechas** (columna `Date`).
5. Crea las medidas: copia cada una desde `dax/measures.dax`
   (`Modelado` > `Nueva medida`). Tabla anfitriona sugerida: `""" + MEASURES_TABLE + """`.
6. Construye las páginas/visuales según `dashboard/dashboard_metadata.json`.

## Opción B — Modelo TMDL (avanzada)

El folder `tmdl/` contiene la definición del modelo semántico en formato TMDL
(`model.tmdl` y `tables/*.tmdl`), con columnas, tipos, particiones Power Query,
relaciones y medidas DAX ya incluidas.

Para usarlo:
- **Tabular Editor 2/3**: abre un modelo y usa *Scripting* o importa los TMDL, o
- **Proyecto Power BI (.pbip)**: copia el contenido de `tmdl/` dentro de
  `<Proyecto>.SemanticModel/definition/` y abre el `.pbip` en Power BI Desktop
  (Vista previa de proyectos PBIP habilitada).

> Nota: las particiones M apuntan a la ruta absoluta de los CSV en `data/`.
> Si mueves los archivos, actualiza la ruta en cada `tables/*.tmdl` o
> parametriza una `DataFolder` en Power Query.

## Actualización de datos

Regenera todo tras editar el Excel ejecutando:

```
python build_powerbi_export.py
```
"""
    with open(os.path.join(OUT_DIR, "README.md"), "w", encoding="utf-8") as f:
        f.write(content)
    print("    - README.md")


# ═════════════════════════════════════════════════════════════
# MAIN
# ═════════════════════════════════════════════════════════════
def main():
    print("=" * 64)
    print(" EXPORTADOR A POWER BI — Transformación Digital gloch.ec")
    print("=" * 64)

    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"No se encontró {EXCEL_PATH}")

    for d in (OUT_DIR, DATA_DIR, DAX_DIR, TMDL_DIR, TABLES_DIR, DASH_DIR):
        os.makedirs(d, exist_ok=True)

    # 1-2) Datos
    tables = read_tables()
    tables = recompute_derived(tables)
    dims = build_dimensions(tables)

    all_tables = {}
    all_tables.update(tables)
    all_tables.update(dims)

    # 3) Genera datos
    write_csvs(all_tables)
    # 4) Genera DAX
    write_dax()
    # 5) Genera TMDL
    write_tmdl(all_tables)
    # 6) Genera metadatos del dashboard
    write_dashboard_metadata()
    # 7) README
    write_readme(all_tables)

    print("=" * 64)
    print(f" LISTO. Paquete generado en: {OUT_DIR}")
    print("=" * 64)


if __name__ == "__main__":
    main()
